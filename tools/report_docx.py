"""
report_docx.py -- 報告書の木（Markdown から作った要素の木）を Word（.docx）に組み立てる部品。

export_report.py から使う部品。直接は実行しない。

課長が自分で手直しできるよう、Word の標準の書式（見出し・箇条書き・表のスタイル）を
そのまま使う。python-docx の既定テンプレートから作り、和文フォントは各スタイルの
eastAsia に明示する（既定テーマのままだと和文が明朝などになるため）。
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Length, Mm, Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

_LATIN_FONT = "Yu Gothic"
_EAST_ASIA_FONT = "游ゴシック"
_CODE_FONT = "Consolas"
_BODY_SIZE = Pt(10.5)
_CODE_SIZE = Pt(9)
_HEADER_FILL = "F6F8FA"
_RULE_COLOR = "D1D9E0"
_MAX_IMAGE_HEIGHT = Mm(150)
_QUOTE_INDENT = Mm(6)

_PAGE = {"width": Mm(210), "height": Mm(297)}
_MARGINS = {"left": Mm(16), "right": Mm(16), "top": Mm(18), "bottom": Mm(20)}

#: Markdown の見出しの段と、Word のスタイル名の対応（h4 以下は Heading 3）。
_HEADING_STYLES = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2"}
_HEADINGS = ("h1", "h2", "h3", "h4", "h5", "h6")
_BLOCK_TAGS = frozenset(
    {"p", "ul", "ol", "pre", "table", "blockquote", "hr", "div", *_HEADINGS}
)
_FONT_STYLES = (
    "Normal",
    "Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "List Bullet",
    "List Bullet 2",
    "List Bullet 3",
    "List Number",
    "List Number 2",
    "List Number 3",
)
_MAX_LIST_DEPTH = 3

#: 段落の書式（w:pPr）の中で w:shd より後ろに来る要素（スキーマの順序を守るため）。
_PPR_AFTER_SHADING = (
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)


# rpr_owner は w:style 要素（CT_Style）か w:r 要素（CT_R）。python-docx はこの二つに
# 共通の型（get_or_add_rPr を持つ型）を公開していないので Any で受ける。
def _set_fonts(rpr_owner: Any, latin: str, east_asia: str) -> None:
    """スタイルか run の rFonts に欧文・和文の書体を明示し、テーマ指定を外す。"""
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(key), None)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def _setup_document(doc: DocxDocument, title: str) -> None:
    """用紙・余白・フォント・フッターのページ番号・文書の題名を整える。"""
    section = doc.sections[0]
    section.page_width, section.page_height = _PAGE["width"], _PAGE["height"]
    section.left_margin, section.right_margin = _MARGINS["left"], _MARGINS["right"]
    section.top_margin, section.bottom_margin = _MARGINS["top"], _MARGINS["bottom"]
    for name in _FONT_STYLES:
        _set_fonts(doc.styles[name].element, _LATIN_FONT, _EAST_ASIA_FONT)
    doc.styles["Normal"].font.size = _BODY_SIZE
    _add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = title


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    """段落に Word のフィールド（PAGE など）を差し込む。"""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        if kind is None:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)


def _add_page_number(paragraph: Paragraph) -> None:
    """フッターの中央に「ページ番号 / 総ページ数」を置く。"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(paragraph, "PAGE")
    paragraph.add_run(" / ")
    _add_field(paragraph, "NUMPAGES")


def add_inline(
    paragraph: Paragraph,
    el: ET.Element,
    *,
    bold: bool = False,
    code: bool = False,
    skip: frozenset[str] = frozenset(),
) -> None:
    """要素の中身を run として段落に足す（*skip* の子要素は飛ばす）。"""
    _add_text(paragraph, el.text, bold, code)
    for child in el:
        if child.tag not in skip:
            _add_child(paragraph, child, bold, code)
        tail = child.tail or ""
        _add_text(paragraph, tail.lstrip() if child.tag == "br" else tail, bold, code)


def _add_child(paragraph: Paragraph, child: ET.Element, bold: bool, code: bool) -> None:
    """段落の中の子要素一つを run にする。画像は別に置くので飛ばす。"""
    if child.tag == "br":
        paragraph.add_run().add_break()
    elif child.tag == "img":
        return
    elif child.tag in ("em", "i"):
        for run in _runs_of(paragraph, child, bold, code):
            run.italic = True
    else:  # strong・code・a など
        add_inline(
            paragraph,
            child,
            bold=bold or child.tag in ("strong", "b"),
            code=code or child.tag == "code",
        )


def _runs_of(paragraph: Paragraph, el: ET.Element, bold: bool, code: bool) -> list[Run]:
    """*el* の中身を足し、そのとき増えた run の一覧を返す。"""
    before = len(paragraph.runs)
    add_inline(paragraph, el, bold=bold, code=code)
    return paragraph.runs[before:]


def _add_text(paragraph: Paragraph, text: str | None, bold: bool, code: bool) -> None:
    """文字を一つの run として足す。空なら何もしない。

    Markdown の段落内の改行（ソフト改行）は、表示上は空白なので空白にする
    （python-docx は改行文字を Word の改行に変えてしまうため）。
    """
    if not text:
        return
    run = paragraph.add_run(text.replace("\n", " "))
    run.bold = bold or None
    if code:
        _set_fonts(run._r, _CODE_FONT, _EAST_ASIA_FONT)
        run.font.size = _CODE_SIZE


def _new_numbering(doc: DocxDocument, style_name: str, start: int) -> int:
    """番号付きスタイルの番号を *start* から振り直す新しい番号定義を作り、その番号を返す。"""
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles[style_name].element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(style_num_id).abstractNumId.val
    num = numbering.add_num(abstract_id)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(start)
    return num.numId


def _apply_numbering(paragraph: Paragraph, num_id: int) -> None:
    """段落に番号定義 *num_id* を当てる。"""
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def _length(value: Length | None) -> Length:
    """python-docx が None もあり得るとする長さを、値として取り出す。

    既定テンプレートから作った文書では用紙・余白・図の寸法は必ず入っているので、
    None なら組み立ての前提が崩れている。黙って進めずに止める。
    """
    if value is None:
        raise ValueError("用紙・余白・図の寸法が文書に入っていません")
    return value


class _Builder:
    """要素の木を Word の段落・表・図に変える。"""

    def __init__(self, doc: DocxDocument, images: Mapping[str, Path]) -> None:
        self.doc = doc
        self.images = images
        section = doc.sections[0]
        self.width = _length(section.page_width) - (
            _length(section.left_margin) + _length(section.right_margin)
        )

    def image(self, src: str) -> None:
        """確認済みの画像を、本文幅と高さの上限に収まるよう中央に置く。"""
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = paragraph.add_run().add_picture(str(self.images[src]))
        width, height = _length(shape.width), _length(shape.height)
        scale = min(self.width / width, _MAX_IMAGE_HEIGHT / height, 1.0)
        shape.width, shape.height = Emu(int(width * scale)), Emu(int(height * scale))

    def paragraph(self, el: ET.Element, style: str | None = None) -> None:
        """段落を置く。中の画像は段落の後ろに図として並べる。"""
        if "".join(el.itertext()).strip() or el.find("br") is not None:
            add_inline(self.doc.add_paragraph(style=style), el)
        for img in el.iter("img"):
            self.image(img.get("src", ""))

    def heading(self, el: ET.Element) -> None:
        """見出しを Word の Title / Heading スタイルで置く。"""
        add_inline(
            self.doc.add_paragraph(style=_HEADING_STYLES.get(el.tag, "Heading 3")), el
        )

    def table(self, el: ET.Element) -> None:
        """表を Table Grid で置く。見出し行は太字・網掛け・ページごとに繰り返す。"""
        rows = [
            row
            for part in el
            for row in (part if part.tag in ("thead", "tbody") else [part])
        ]
        ncol = max(len(row) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncol, style="Table Grid")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_el, row in zip(rows, table.rows):
            is_header = any(cell.tag == "th" for cell in row_el)
            if is_header:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell_el, cell in zip(row_el, row.cells):
                add_inline(cell.paragraphs[0], cell_el, bold=is_header)
                if is_header:
                    _shade(cell, _HEADER_FILL)
        self.doc.add_paragraph()

    def bullet_list(self, el: ET.Element, depth: int = 1) -> None:
        """箇条書き・番号付きを List Bullet / List Number（入れ子は 2・3）で置く。"""
        base = "List Number" if el.tag == "ol" else "List Bullet"
        style = base if depth == 1 else f"{base} {min(depth, _MAX_LIST_DEPTH)}"
        num_id = None
        if el.tag == "ol":
            num_id = _new_numbering(self.doc, style, int(el.get("start", "1")))
        for li in (child for child in el if child.tag == "li"):
            paragraph = self.doc.add_paragraph(style=style)
            if num_id is not None:
                _apply_numbering(paragraph, num_id)
            self.list_item(li, paragraph, depth)

    def list_item(self, li: ET.Element, paragraph: Paragraph, depth: int) -> None:
        """箇条書きの一項目の中身を置く。最初の段落は項目の行に入れる。"""
        if (li.text or "").strip() or any(c.tag not in _BLOCK_TAGS for c in li):
            add_inline(paragraph, li, skip=_BLOCK_TAGS)
        for child in li:
            if child.tag in ("ul", "ol"):
                self.bullet_list(child, depth + 1)
            elif child.tag == "p" and not paragraph.text.strip():
                add_inline(paragraph, child)
            elif child.tag == "p":
                self.paragraph(child, style="List Continue")
            elif child.tag in _BLOCK_TAGS:
                self.block(child)

    def code_block(self, el: ET.Element) -> None:
        """コードブロックを、等幅・網掛けの段落として置く。"""
        paragraph = self.doc.add_paragraph()
        _shade_paragraph(paragraph, _HEADER_FILL)
        for index, line in enumerate("".join(el.itertext()).rstrip("\n").split("\n")):
            if index:
                paragraph.runs[-1].add_break()
            _add_text(paragraph, line or " ", bold=False, code=True)

    def quote(self, el: ET.Element) -> None:
        """引用を、字下げと左の縦線の付いた段落として置く。"""
        before = len(self.doc.paragraphs)
        for child in el:
            self.block(child)
        for paragraph in self.doc.paragraphs[before:]:
            paragraph.paragraph_format.left_indent = _QUOTE_INDENT
            _border(paragraph, "left")

    def rule(self) -> None:
        """区切り線を、空の段落の下罫線として置く。"""
        _border(self.doc.add_paragraph(), "bottom")

    def block(self, el: ET.Element) -> None:
        """ブロック要素一つを文書に足す。"""
        handlers = {
            "ul": self.bullet_list,
            "ol": self.bullet_list,
            "table": self.table,
            "pre": self.code_block,
            "blockquote": self.quote,
        }
        if el.tag in _HEADINGS:
            self.heading(el)
        elif el.tag == "hr":
            self.rule()
        elif el.tag in handlers:
            handlers[el.tag](el)
        else:
            self.paragraph(el)


def _shade(cell: _Cell, fill: str) -> None:
    """表のセルに網掛けの色を付ける。"""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _shade_paragraph(paragraph: Paragraph, fill: str) -> None:
    """段落に網掛けの色を付ける。"""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().insert_element_before(shading, *_PPR_AFTER_SHADING)


def _border(paragraph: Paragraph, side: str) -> None:
    """段落の *side*（bottom / left）に罫線を引く。"""
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.insert_element_before(borders, "w:shd", *_PPR_AFTER_SHADING)
    line = OxmlElement(f"w:{side}")
    for key, value in (
        ("val", "single"),
        ("sz", "12"),
        ("space", "4"),
        ("color", _RULE_COLOR),
    ):
        line.set(qn(f"w:{key}"), value)
    borders.append(line)


def build_docx(root: ET.Element, images: Mapping[str, Path], title: str) -> bytes:
    """要素の木 *root* から Word 文書を組み立て、そのバイト列を返す。

    *images* は報告書中の画像の src から、確認済みの実在ファイルへの対応。
    *title* は文書のプロパティの題名に入れる。
    """
    doc = Document()
    _setup_document(doc, title)
    builder = _Builder(doc, images)
    for el in root:
        builder.block(el)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
